#!/usr/bin/env python3
"""
CV generator. Output is a .docx you can open in Pages or Word.
Layout follows a classic US "General resume": Inter, name on the left,
uppercase section headings, dates on the right.

Facts are NOT typed here. Everything is read from one source file:
    cv_data.md  →  the block between <!-- CV:START --> and <!-- CV:END -->

Usage:
    python3 cv_build.py                  # list available versions
    python3 cv_build.py pm               # build a version, no photo (US, UK, Ireland, NL...)
    python3 cv_build.py pm --photo       # photo on the right of the header (DACH, Balkans)

A photo is also added when the version in cv_data.md has `photo: yes`.
For a new job ad: add a `### @version companyname` block to cv_data.md.
Output goes to ./output/ as CV_<version>.docx.
"""

import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
SOURCE = HERE / "cv_data.md"
OUTPUT_DIR = HERE / "output"
PHOTO = HERE / "headshot.png"

FONT = "Inter"
TEXT_SIZE = 9.5
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x40, 0x40, 0x40)
RIGHT_TAB = Inches(7.5)


# ─────────────────────────────────────────────────────────────────────────
#  READING cv_data.md
# ─────────────────────────────────────────────────────────────────────────
def load():
    """Return blocks as {'contact': {...}, 'version': {...}, ...}."""
    text = SOURCE.read_text(encoding="utf-8")
    m = re.search(r"^<!-- CV:START -->$(.*?)^<!-- CV:END -->$", text, re.S | re.M)
    if not m:
        sys.exit(f"No CV:START/CV:END block found in {SOURCE}")

    blocks = {}
    current = None
    for line in m.group(1).split("\n"):
        header = re.match(r"^###\s+@(\w+)(?:\s+(\S+))?\s*$", line)
        if header:
            kind, name = header.group(1), header.group(2)
            current = {"_kind": kind, "_name": name, "bullets": [], "details": [], "lines": []}
            blocks.setdefault(kind, {})[name or kind] = current
            continue
        if current is None or not line.strip():
            continue
        if line.startswith("- "):
            current["bullets"].append(line[2:].strip())
        elif ":" in line and re.match(r"^[a-z]+:", line):
            key, value = line.split(":", 1)
            if key == "detail":
                current["details"].append(value.strip())
            else:
                current[key] = value.strip()
        else:
            current["lines"].append(line.strip())
    return blocks


# ─────────────────────────────────────────────────────────────────────────
#  DOCX HELPERS
# ─────────────────────────────────────────────────────────────────────────
def clean(text):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def run(p, text, bold=False, italic=False, size=TEXT_SIZE, color=BLACK):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return r


def paragraph(doc, before=0, after=0, right_tab=False, style=None, existing=None, keep_next=False):
    p = existing or (doc.add_paragraph(style=style) if style else doc.add_paragraph())
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), 1.0
    f.keep_with_next = keep_next
    if right_tab:
        f.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    return p


def left_right(doc, left, right, bold=True, before=0):
    p = paragraph(doc, before=before, right_tab=True, keep_next=True)
    run(p, left, bold=bold)
    if right:
        run(p, "\t" + right, bold=bold)
    return p


def section(doc, title):
    run(paragraph(doc, before=10, after=4, keep_next=True), title, bold=True, size=11)


def split_company(company):
    if ", " in company:
        name, city = company.rsplit(", ", 1)
        return name, city
    return company, ""


def split_education(line):
    parts = [d.strip() for d in line.split("·") if d.strip()]
    title = parts[0]
    date = next((d for d in parts[1:] if re.search(r"\d{4}", d) and not d.startswith("GPA")), "")
    rest = [d for d in parts[1:] if d != date]
    return title, date, rest


def borderless_table(table):
    """Full width table with no borders, cell margins or indent."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(7.5 * 1440)))
    width.set(qn("w:type"), "dxa")
    tblPr.append(width)
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right", "top", "bottom"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), "0")
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    tblPr.append(margins)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    tblPr.append(indent)


# ─────────────────────────────────────────────────────────────────────────
#  BUILDING THE CV
# ─────────────────────────────────────────────────────────────────────────
def build(version_name, photo=False):
    blocks = load()
    versions = blocks.get("version", {})
    if version_name not in versions:
        sys.exit(f"No version '{version_name}'. Available: {', '.join(sorted(versions))}")
    v = versions[version_name]
    contact = next(iter(blocks["contact"].values()))
    photo = photo or v.get("photo", "").lower() in ("yes", "true")

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, side, Inches(0.5))
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(TEXT_SIZE)

    # ── Header: name and contact on the left, optional photo on the right ──
    header, first = doc, None
    if photo:
        if not PHOTO.exists():
            sys.exit(f"Photo not found: {PHOTO}")
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        borderless_table(table)
        left, right = table.rows[0].cells
        for column, cell, width in ((table.columns[0], left, 6.4), (table.columns[1], right, 1.1)):
            column.width = cell.width = Inches(width)
        img = right.paragraphs[0]
        img.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        img.add_run().add_picture(str(PHOTO), width=Inches(0.95))
        header, first = left, left.paragraphs[0]

    run(paragraph(header, existing=first), contact.get("name", "").title(), bold=True, size=18)
    if v.get("title"):
        run(paragraph(header, after=2), v["title"].replace("  ·  ", " | "), color=GREY)
    city = v.get("city") or contact.get("city", "")
    run(paragraph(header), f"{city} | Phone: {contact.get('phone', '')}")
    run(paragraph(header), f"Email: {contact.get('email', '')} | LinkedIn: {contact.get('linkedin', '')}")

    # ── Summary ──
    if v.get("summary"):
        section(doc, "SUMMARY")
        run(paragraph(doc), clean(v["summary"]))

    # ── Experience ──
    jobs = [x.strip() for x in v.get("jobs", "").split(",") if x.strip()]
    if jobs:
        section(doc, "PROFESSIONAL EXPERIENCE")
        for i, key in enumerate(jobs):
            job = blocks.get("job", {}).get(key)
            if not job:
                print(f"⚠️  no @job {key} block, skipping")
                continue
            left_right(doc, job.get("role", ""), job.get("period", ""), before=0 if i == 0 else 8)
            company, city = split_company(job.get("company", ""))
            left_right(doc, company, city)
            if job.get("subtitle"):
                run(paragraph(doc, keep_next=True), job["subtitle"])
            if job["bullets"]:
                run(paragraph(doc, keep_next=True), "Key Accomplishments:", italic=True)
                for b in job["bullets"]:
                    p = paragraph(doc, style="List Bullet")
                    p.paragraph_format.left_indent = Inches(0.5)
                    p.paragraph_format.first_line_indent = Inches(-0.25)
                    run(p, clean(b))

    # ── Education ──
    items = [x.strip() for x in v.get("education", "").split(",") if x.strip()]
    if items:
        section(doc, "EDUCATION")
        for i, key in enumerate(items):
            e = blocks.get("education", {}).get(key)
            if not e:
                print(f"⚠️  no @education {key} block, skipping")
                continue
            title, date, rest = split_education(e.get("line", ""))
            degree, _, institution = title.partition(" — ")
            left_right(doc, degree, date, before=0 if i == 0 else 6)
            if institution:
                run(paragraph(doc), institution, bold=True)
            details = ([" · ".join(rest)] if rest else []) + e["details"]
            for d in details:
                run(paragraph(doc), clean(d), size=TEXT_SIZE - 0.5, color=GREY)

    # ── Skills ──
    if v.get("skills", "").lower() in ("yes", "true"):
        skills = blocks.get("skills", {}).get("skills")
        if skills:
            section(doc, "ADDITIONAL INFORMATION")
            for line in skills["lines"]:
                p = paragraph(doc, after=2)
                if ":" in line:
                    head, tail = line.split(":", 1)
                    run(p, head + ":", bold=True)
                    run(p, tail)
                else:
                    run(p, line)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / f"CV_{version_name}.docx"
    doc.save(out)
    print(f"✅ {out.name} · {out.stat().st_size // 1024} KB · {'with photo' if photo else 'no photo'}")
    print(f"   {out}")
    print("   Open in Pages or Word, check page breaks, export to PDF.")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args or "--list" in sys.argv or "-l" in sys.argv:
        b = load()
        print("Versions:", ", ".join(sorted(b.get("version", {}))))
        print("Jobs:", ", ".join(sorted(b.get("job", {}))))
        print("Education:", ", ".join(sorted(b.get("education", {}))))
    else:
        build(args[0], photo="--photo" in sys.argv)
